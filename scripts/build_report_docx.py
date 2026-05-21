from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "课程设计报告.md"
TARGET = ROOT / "docs" / "课程设计报告.docx"


def add_markdown(document: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    idx = 0
    in_code = False
    while idx < len(lines):
        raw = lines[idx]
        line = raw.rstrip()
        if not line:
            document.add_paragraph()
            idx += 1
            continue
        if line.startswith("```"):
            in_code = not in_code
            idx += 1
            continue
        if in_code:
            paragraph = document.add_paragraph()
            paragraph.add_run(line).font.name = "Consolas"
            idx += 1
            continue
        if line.startswith("!["):
            add_image(document, line)
            idx += 1
            continue
        if is_table_start(lines, idx):
            idx = add_table(document, lines, idx)
            continue
        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            document.add_heading(line[4:].strip(), level=3)
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(clean_inline(line[2:].strip()))
        else:
            document.add_paragraph(clean_inline(line))
        idx += 1


def clean_inline(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text


def add_image(document: Document, line: str) -> None:
    match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line.strip())
    if not match:
        document.add_paragraph(clean_inline(line))
        return
    caption, raw_path = match.groups()
    image_path = Path(raw_path.strip())
    if not image_path.is_absolute():
        image_path = (SOURCE.parent / image_path).resolve()
        if not image_path.exists():
            image_path = (ROOT / raw_path.strip()).resolve()
    if image_path.exists():
        document.add_picture(str(image_path), width=Inches(6.0))
        if caption:
            document.add_paragraph(caption)
    else:
        document.add_paragraph(f"{caption}：{raw_path}")


def is_table_start(lines: list[str], idx: int) -> bool:
    if idx + 1 >= len(lines):
        return False
    return lines[idx].lstrip().startswith("|") and re.match(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", lines[idx + 1])


def parse_table_row(line: str) -> list[str]:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return [clean_inline(cell) for cell in cells]


def add_table(document: Document, lines: list[str], idx: int) -> int:
    headers = parse_table_row(lines[idx])
    rows: list[list[str]] = []
    idx += 2
    while idx < len(lines) and lines[idx].lstrip().startswith("|"):
        rows.append(parse_table_row(lines[idx]))
        idx += 1

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for col_idx, header in enumerate(headers):
        table.rows[0].cells[col_idx].text = header
    for row in rows:
        cells = table.add_row().cells
        for col_idx, value in enumerate(row[: len(headers)]):
            cells[col_idx].text = value
    return idx


def main() -> None:
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"].font.size = Pt(10.5)
    add_markdown(document, SOURCE.read_text(encoding="utf-8"))
    document.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
